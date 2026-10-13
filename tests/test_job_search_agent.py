import csv
import json
import tempfile
import unittest
from argparse import Namespace
from pathlib import Path
from unittest.mock import Mock, patch

from docx import Document

import job_search_agent as agent


class JobSearchAgentTests(unittest.TestCase):
    def test_verified_smb_requires_explicit_size_at_or_below_500(self):
        self.assertTrue(
            agent._is_verified_smb(
                "51-200 employees", "Acme has 51-200 employees."
            )
        )
        self.assertTrue(
            agent._is_verified_smb("500 employees", "Acme has 500 employees.")
        )
        self.assertFalse(
            agent._is_verified_smb(
                "201-1,000 employees", "Acme has 201-1,000 employees."
            )
        )
        self.assertFalse(
            agent._is_verified_smb(
                "501 employees", "Acme has 501 employees."
            )
        )
        self.assertFalse(
            agent._is_verified_smb("Not verified", "No employee count found.")
        )
        self.assertFalse(
            agent._is_verified_smb("51-200 employees", "No employee count found.")
        )

    def test_extract_search_queries_validates_and_deduplicates(self):
        with patch.object(
            agent,
            "_ollama_generate",
            return_value=json.dumps({"search_queries": ["Data analyst", "Data analyst", "SQL"]}),
        ):
            queries = agent.extract_search_queries("resume text", "llama3.2:3b", "http://localhost:11434")
        self.assertEqual(queries, ["Data analyst", "SQL"])

    def test_extract_search_queries_includes_target_industry(self):
        with patch.object(
            agent,
            "_ollama_generate",
            return_value=json.dumps({"search_queries": ["data analyst"]}),
        ) as generate:
            agent.extract_search_queries(
                "resume text", "llama3.2:3b", "http://localhost:11434", "healthcare"
            )
        self.assertIn("healthcare industry", generate.call_args.args[0])

    def test_extract_search_queries_any_industry_does_not_restrict_sector(self):
        with patch.object(
            agent,
            "_ollama_generate",
            return_value=json.dumps({"search_queries": ["data analyst"]}),
        ) as generate:
            agent.extract_search_queries(
                "resume text", "llama3.2:3b", "http://localhost:11434", "any"
            )
        prompt = generate.call_args.args[0]
        self.assertIn("Search across any industry", prompt)
        self.assertNotIn("Focus the search terms on the any industry", prompt)

    def test_assess_posting_rejects_invalid_requirement_assessment(self):
        posting = agent.Posting("1", "Analyst", "Acme", "Remote", "SQL", "https://example.com")
        with patch.object(
            agent,
            "_ollama_generate",
            side_effect=[
                json.dumps(
                    {
                        "is_individual_posting": True,
                        "industry_match": True,
                        "location_match": "match",
                        "posting_location": "Remote",
                        "requirements": [
                            {
                                "criterion": "SQL",
                                "priority": "essential",
                                "status": "met",
                                "evidence": "SQL",
                            }
                        ],
                        "summary": "Strong match",
                    }
                ),
                json.dumps(
                {
                    "is_individual_posting": True,
                    "industry_match": True,
                    "location_match": "match",
                    "posting_location": "Remote",
                    "requirements": [
                        {
                            "criterion": "SQL",
                            "priority": "essential",
                            "status": "met",
                            "evidence": "SQL",
                        }
                    ],
                    "summary": "Strong match",
                }
                ),
            ],
        ):
            with self.assertRaisesRegex(RuntimeError, "invalid requirement"):
                agent.assess_posting("resume", posting, "llama3.2:3b", "http://localhost:11434")

    def test_assess_posting_retries_incomplete_model_assessment_once(self):
        posting = agent.Posting(
            "1",
            "Research Engineer",
            "Waabi",
            "Toronto",
            "Develop research prototypes using machine learning.",
            "https://jobs.example/research-engineer",
        )
        incomplete = {
            "is_individual_posting": True,
            "requirements": [],
            "summary": "Research role identified.",
        }
        valid = {
            "is_individual_posting": True,
            "industry_match": True,
            "location_match": "match",
            "posting_location": "Toronto, Canada",
            "employer_name": "Waabi",
            "requirements": [
                {
                    "criterion": "Machine learning research",
                    "priority": "preferred",
                    "status": "met",
                    "evidence": "Built machine learning prototypes.",
                }
            ],
            "summary": "The role aligns with the candidate's research experience.",
        }
        with patch.object(
            agent,
            "_ollama_generate",
            side_effect=[json.dumps(incomplete), json.dumps(valid)],
        ) as generate:
            assessment = agent.assess_posting(
                "Built machine learning prototypes.",
                posting,
                "llama3.2:3b",
                "http://localhost:11434",
                "autonomous vehicles",
            )

        self.assertTrue(assessment.qualified)
        self.assertEqual(assessment.posting.company, "Waabi")
        self.assertEqual(assessment.score, 100)
        self.assertEqual(generate.call_count, 2)
        self.assertIn("Return only the corrected JSON object", generate.call_args.args[0])

    def test_create_report_includes_all_postings_and_marks_qualification(self):
        qualified = agent.Posting(
            "1", "Analyst", "Acme", "Remote", "Build useful reports.", "https://example.com"
        )
        rejected = agent.Posting("2", "Engineer", "Other", "Remote", "", "https://example.org")
        assessments = [
            agent.Assessment(
                qualified, 80, True, "Good evidence", ["SQL"], [], 100, 33
            ),
            agent.Assessment(rejected, 95, False, "Missing requirements", [], ["Python"]),
        ]
        with tempfile.TemporaryDirectory() as directory:
            output = Path(directory) / "jobs.docx"
            agent.create_report(assessments, output, "Remote", ["analyst"], 60, "healthcare")
            document = Document(output)
        content = "\n".join(paragraph.text for paragraph in document.paragraphs)
        self.assertIn("Analyst", content)
        self.assertIn("Build useful reports.", content)
        self.assertIn("Industry focus: healthcare", content)
        self.assertIn("ATS-style match score: 80/100", content)
        self.assertIn("Required criteria coverage: 100%", content)
        self.assertIn("Preferred criteria coverage: 33%", content)
        self.assertIn("Link check: Not checked", content)
        self.assertIn("Engineer", content)
        self.assertIn("Assessment: Not qualified / needs review", content)
        self.assertTrue(
            any(
                rel.is_external and rel.target_ref == "https://example.com"
                for rel in document.part.rels.values()
            )
        )

    def test_create_csv_report_exports_all_jobs_with_details(self):
        passing = agent.Posting(
            "1",
            "Analyst",
            "Acme",
            "Toronto, Ontario, Canada",
            "Build reports",
            "https://example.com/job",
            80000,
            95000,
            True,
            "HTTP 200 - reachable",
            "https://careers.example.com/jobs/123",
            search_date="2026-10-10",
        )
        failing = agent.Posting(
            "2", "Engineer", "Other", "Remote", "", "https://example.org/job"
        )
        assessments = [
            agent.Assessment(
                passing, 80, True, "Good fit", ["SQL", "Excel"], ["Cloud"], 100, 33
            ),
            agent.Assessment(failing, 99, False, "Not qualified", [], []),
        ]
        with tempfile.TemporaryDirectory() as directory:
            output = Path(directory) / "nested" / "jobs.csv"
            agent.create_csv_report(assessments, output, min_score=60)
            with output.open(encoding="utf-8-sig", newline="") as csv_file:
                rows = list(csv.DictReader(csv_file))
        self.assertEqual(len(rows), 2)
        self.assertEqual(rows[0]["title"], "Analyst")
        self.assertEqual(rows[0]["location"], "Toronto, Ontario, Canada")
        self.assertEqual(rows[0]["location_match"], "unclear")
        self.assertEqual(rows[0]["search_date"], "2026-10-10")
        self.assertEqual(rows[0]["company_size"], "Not verified")
        self.assertEqual(rows[0]["hiring_status"], "Not verified")
        self.assertEqual(rows[0]["salary_max"], "95000")
        self.assertEqual(rows[0]["required_coverage"], "100")
        self.assertEqual(rows[0]["preferred_coverage"], "33")
        self.assertEqual(rows[0]["matching_skills"], "SQL; Excel")
        self.assertEqual(rows[0]["url"], "https://example.com/job")
        self.assertEqual(
            rows[0]["apply_url"], "https://careers.example.com/jobs/123"
        )
        self.assertEqual(rows[0]["link_verified"], "True")
        self.assertEqual(rows[0]["link_status"], "HTTP 200 - reachable")
        self.assertEqual(rows[0]["final_url"], "https://careers.example.com/jobs/123")
        rejected_row = next(row for row in rows if row["title"] == "Engineer")
        self.assertEqual(rejected_row["qualified"], "False")
        self.assertEqual(rejected_row["meets_threshold"], "False")

    def test_create_csv_report_wraps_long_titles_without_truncating(self):
        title = ("Senior Research Engineer " * 6).strip()
        assessment = agent.Assessment(
            agent.Posting(
                "1", title, "Acme", "Canada", "", "https://example.com/job"
            ),
            80,
            True,
            "Strong match",
            [],
            [],
        )
        with tempfile.TemporaryDirectory() as directory:
            output = Path(directory) / "jobs.csv"
            agent.create_csv_report([assessment], output, min_score=60)
            with output.open(encoding="utf-8-sig", newline="") as csv_file:
                rows = list(csv.DictReader(csv_file))

        wrapped_title = rows[0]["title"]
        self.assertIn("\n", wrapped_title)
        self.assertEqual(" ".join(wrapped_title.splitlines()), title)
        self.assertTrue(
            all(len(line) <= 100 for line in wrapped_title.splitlines())
        )

    def test_verify_employer_facts_records_size_and_individual_opening_evidence(self):
        assessment = agent.Assessment(
            agent.Posting(
                "1",
                "Data Engineer",
                "Acme",
                "Remote",
                "Build data pipelines.",
                "https://acme.example/jobs/data-engineer",
            ),
            80,
            True,
            "Strong match",
            ["Python"],
            [],
        )
        search = Mock()
        search.results.side_effect = [
            [
                {
                    "title": "Acme company size",
                    "link": "https://directory.example/acme",
                    "snippet": "Acme has 51-200 employees.",
                }
            ],
            [
                {
                    "title": "Data Engineer at Acme",
                    "link": "https://acme.example/jobs/data-engineer",
                    "snippet": "Acme is hiring a Data Engineer. Apply now.",
                }
            ],
        ]
        with patch(
            "langchain_community.utilities.DuckDuckGoSearchAPIWrapper",
            return_value=search,
        ):
            result = agent.verify_employer_facts([assessment])

        posting = result[0].posting
        self.assertEqual(posting.company_size, "51-200 employees")
        self.assertEqual(posting.company_size_url, "https://directory.example/acme")
        self.assertIn("51-200 employees", posting.company_size_evidence)
        self.assertEqual(
            posting.hiring_status,
            "Individual posting found in search; active status not confirmed",
        )
        self.assertEqual(
            posting.hiring_evidence_url,
            "https://acme.example/jobs/data-engineer",
        )
        self.assertEqual(search.results.call_count, 2)

    def test_verify_employer_facts_keeps_unknown_when_employer_is_not_identified(self):
        assessment = agent.Assessment(
            agent.Posting(
                "1", "Analyst", "See posting", "Remote", "", "https://example.com/job"
            ),
            50,
            False,
            "Unknown",
            [],
            [],
        )
        with patch("langchain_community.utilities.DuckDuckGoSearchAPIWrapper") as search:
            result = agent.verify_employer_facts([assessment])
        self.assertEqual(result[0].posting.company_size, "Not verified")
        self.assertEqual(result[0].posting.hiring_status, "Not verified")
        search.assert_not_called()

    def test_link_verification_visits_url_and_records_redirect(self):
        posting = agent.Posting(
            "1", "Analyst", "Acme", "Remote", "Analyze data", "https://example.com/job"
        )

        class FakeResponse:
            def __enter__(self):
                return self

            def __exit__(self, *_args):
                return None

            def getcode(self):
                return 200

            def geturl(self):
                return "https://careers.example.com/jobs/123"

            def read(self, _size):
                return b"job posting"

        class FakeOpener:
            def open(self, request, timeout):
                self.request = request
                self.timeout = timeout
                return FakeResponse()

        fake_opener = FakeOpener()
        with (
            patch.object(agent, "_validate_public_url"),
            patch.object(agent, "build_opener", return_value=fake_opener),
        ):
            verified = agent.verify_posting_link(posting)
        self.assertTrue(verified.link_verified)
        self.assertEqual(verified.link_status, "HTTP 200 - reachable")
        self.assertEqual(verified.final_url, "https://careers.example.com/jobs/123")
        self.assertEqual(fake_opener.request.full_url, posting.url)
        self.assertEqual(fake_opener.timeout, 12)

    def test_link_verification_reports_http_error(self):
        posting = agent.Posting(
            "1", "Analyst", "Acme", "Remote", "Analyze data", "https://example.com/job"
        )
        error = agent.HTTPError(
            posting.url, 404, "Not Found", {}, None
        )
        fake_opener = Mock()
        fake_opener.open.side_effect = error
        with (
            patch.object(agent, "_validate_public_url"),
            patch.object(agent, "build_opener", return_value=fake_opener),
        ):
            verified = agent.verify_posting_link(posting)
        self.assertFalse(verified.link_verified)
        self.assertIn("HTTP 404", verified.link_status)

    def test_link_verification_blocks_local_hosts_before_request(self):
        posting = agent.Posting(
            "1", "Analyst", "Acme", "Remote", "Analyze data", "http://127.0.0.1/admin"
        )
        with patch.object(agent, "build_opener") as build_opener:
            verified = agent.verify_posting_link(posting)
        self.assertFalse(verified.link_verified)
        self.assertIn("public IP", verified.link_status)
        build_opener.assert_not_called()

    def test_fetch_jobs_keeps_individual_job_results_without_model_extraction(self):
        search_wrapper = Mock()
        search_wrapper.results.return_value = [
            {
                "title": "Data Analyst - Acme",
                "link": "https://example.com/job",
                "snippet": "Acme analyst role. Analyze data.",
            },
            {
                "title": "Duplicate result",
                "link": "https://example.com/job#apply",
                "snippet": "Duplicate.",
            },
        ]
        with (
            patch(
                "langchain_community.utilities.DuckDuckGoSearchAPIWrapper",
                return_value=search_wrapper,
            ) as search_wrapper_class,
            patch.object(agent, "_ollama_generate") as generate,
        ):
            jobs = agent.fetch_jobs(
                ["analyst"],
                "Remote",
                5,
                "healthcare",
            )
        self.assertEqual(len(jobs), 1)
        self.assertEqual(jobs[0].title, "Data Analyst - Acme")
        self.assertEqual(jobs[0].company, "See posting")
        self.assertEqual(jobs[0].location, "Remote")
        self.assertEqual(jobs[0].description, "Acme analyst role. Analyze data.")
        self.assertEqual(jobs[0].url, "https://example.com/job")
        self.assertEqual(
            jobs[0].search_date,
            agent.datetime.now().astimezone().date().isoformat(),
        )
        self.assertIsNone(jobs[0].salary_min)
        generate.assert_not_called()
        search_wrapper_class.assert_called_once_with(
            backend="bing,brave,google,mojeek,yahoo"
        )
        self.assertEqual(
            [call.args for call in search_wrapper.results.call_args_list],
            [
                (
                    "analyst job opening apply Remote healthcare",
                ),
                (
                    "site:indeed.com/viewjob analyst job opening apply Remote healthcare",
                ),
                (
                    'analyst job opening apply Remote healthcare startup OR "small business" OR SMB',
                ),
            ],
        )
        self.assertTrue(
            all(
                call.kwargs == {"max_results": 5}
                for call in search_wrapper.results.call_args_list
            )
        )

    def test_fetch_jobs_discards_indeed_search_page_but_keeps_direct_job(self):
        search_wrapper = Mock()
        search_wrapper.results.return_value = [
            {
                "title": "$35K-$139K Aerospace Jobs in Canada | Indeed",
                "link": "https://ca.indeed.com/Aerospace-jobs-in-Canada",
                "snippet": "3,507 Aerospace jobs available in Canada. Apply to ... and more!",
            },
            {
                "title": "Avionics Engineer",
                "link": "https://ca.indeed.com/viewjob?jk=abc123",
                "snippet": "Acme is hiring an Avionics Engineer. Apply now.",
            },
            {
                "title": "Search Aerospace jobs",
                "link": "https://www.linkedin.com/jobs/search/?keywords=aerospace",
                "snippet": "Browse aerospace jobs.",
            },
            {
                "title": "Avionics Engineer - Acme Careers",
                "link": "https://careers.acme.example/jobs/avionics-engineer-123",
                "snippet": "Apply for the Avionics Engineer role.",
            },
            {
                "title": "Systems Engineer - Contoso",
                "link": "https://careers.contoso.example/opening/8f12",
                "snippet": "Contoso is hiring a Systems Engineer.",
            },
            {
                "title": "Engineer jobs",
                "link": "https://careers.contoso.example/search?query=engineer",
                "snippet": "Browse open Engineer roles.",
            },
            {
                "title": "Aerospace jobs in Canada",
                "link": "https://jobs.board.example/jobs-in-canada",
                "snippet": "Browse 3,507 aerospace jobs in Canada.",
            },
        ]
        with patch(
            "langchain_community.utilities.DuckDuckGoSearchAPIWrapper",
            return_value=search_wrapper,
        ):
            jobs = agent.fetch_jobs(["aerospace engineer"], "Canada", 10)
        self.assertEqual(
            [job.url for job in jobs],
            [
                "https://ca.indeed.com/viewjob?jk=abc123",
                "https://careers.acme.example/jobs/avionics-engineer-123",
                "https://careers.contoso.example/opening/8f12",
            ],
        )

    def test_fetch_jobs_returns_empty_when_search_has_no_results(self):
        with (
            patch(
                "langchain_community.utilities.DuckDuckGoSearchAPIWrapper"
            ) as search_wrapper_class,
        ):
            search_wrapper_class.return_value.results.return_value = []
            jobs = agent.fetch_jobs(
                ["analyst"], "Remote", 5
            )
        self.assertEqual(jobs, [])

    def test_fetch_jobs_treats_backend_no_results_as_empty_and_continues(self):
        search_wrapper = Mock()
        search_wrapper.results.side_effect = [
            RuntimeError("No results found."),
            [
                {
                    "title": "Analyst - Acme",
                    "link": "https://acme.example/jobs/analyst",
                    "snippet": "Acme is hiring an analyst.",
                }
            ],
            [
                {
                    "title": "Analyst - Acme",
                    "link": "https://acme.example/jobs/analyst",
                    "snippet": "Acme is hiring an analyst.",
                }
            ],
        ]
        with patch(
            "langchain_community.utilities.DuckDuckGoSearchAPIWrapper",
            return_value=search_wrapper,
        ):
            jobs = agent.fetch_jobs(["analyst"], "Canada", 5)
        self.assertEqual(
            [job.url for job in jobs], ["https://acme.example/jobs/analyst"]
        )
        self.assertEqual(search_wrapper.results.call_count, 3)

    def test_fetch_jobs_reports_network_errors_with_actionable_context(self):
        search_wrapper = Mock()
        search_wrapper.results.side_effect = RuntimeError(
            "DNS resolution failed for search.example"
        )
        with patch(
            "langchain_community.utilities.DuckDuckGoSearchAPIWrapper",
            return_value=search_wrapper,
        ):
            with self.assertRaisesRegex(
                RuntimeError, "Check your internet/DNS connection"
            ):
                agent.fetch_jobs(["analyst"], "Canada", 5)

    def test_fetch_jobs_any_industry_does_not_add_sector_filter(self):
        search_wrapper = Mock()
        search_wrapper.results.return_value = [
            {
                "title": "Data Analyst",
                "link": "https://careers.example.com/job/data-analyst-1",
                "snippet": "Analyze data across business teams.",
            }
        ]
        with patch(
            "langchain_community.utilities.DuckDuckGoSearchAPIWrapper",
            return_value=search_wrapper,
        ):
            agent.fetch_jobs(["data analyst"], "Remote", 5, "any")
        self.assertEqual(
            [call.args[0] for call in search_wrapper.results.call_args_list],
            [
                "data analyst job opening apply Remote",
                "site:indeed.com/viewjob data analyst job opening apply Remote",
                'data analyst job opening apply Remote startup OR "small business" OR SMB',
            ],
        )
        self.assertTrue(
            all(
                call.kwargs == {"max_results": 5}
                for call in search_wrapper.results.call_args_list
            )
        )

    def test_fetch_jobs_reserves_results_for_startups_and_smbs(self):
        general_results = [
            {
                "title": f"General role {index}",
                "link": f"https://company{index}.example/jobs/role-{index}",
                "snippet": "Hiring for a role.",
            }
            for index in range(1, 5)
        ]
        startup_results = [
            {
                "title": "Startup Data Analyst",
                "link": "https://startup.example/jobs/data-analyst",
                "snippet": "Early-stage startup hiring a data analyst.",
            }
        ]
        search_wrapper = Mock()
        search_wrapper.results.side_effect = [general_results, [], startup_results]
        with patch(
            "langchain_community.utilities.DuckDuckGoSearchAPIWrapper",
            return_value=search_wrapper,
        ):
            jobs = agent.fetch_jobs(["data analyst"], "Remote", 3)
        self.assertEqual(
            [job.url for job in jobs],
            [
                "https://company1.example/jobs/role-1",
                "https://startup.example/jobs/data-analyst",
                "https://company2.example/jobs/role-2",
            ],
        )
        self.assertIn("startup OR", search_wrapper.results.call_args_list[2].args[0])

    def test_fetch_jobs_searches_indeed_for_direct_postings(self):
        search_wrapper = Mock()
        search_wrapper.results.side_effect = [
            [],
            [
                {
                    "title": "Data Analyst - Acme",
                    "link": "https://ca.indeed.com/viewjob?jk=abc123",
                    "snippet": "Acme is hiring a data analyst.",
                }
            ],
            [],
        ]
        with patch(
            "langchain_community.utilities.DuckDuckGoSearchAPIWrapper",
            return_value=search_wrapper,
        ):
            jobs = agent.fetch_jobs(["data analyst"], "Canada", 3)

        self.assertEqual(
            [job.url for job in jobs],
            ["https://ca.indeed.com/viewjob?jk=abc123"],
        )
        self.assertEqual(
            search_wrapper.results.call_args_list[1].args[0],
            "site:indeed.com/viewjob data analyst job opening apply Canada",
        )

    def test_fetch_jobs_smb_only_targets_the_selected_industry(self):
        search_wrapper = Mock()
        search_wrapper.results.return_value = []
        with patch(
            "langchain_community.utilities.DuckDuckGoSearchAPIWrapper",
            return_value=search_wrapper,
        ):
            agent.fetch_jobs(
                ["aerospace engineer"],
                "Canada",
                5,
                "aerospace",
                smb_only=True,
            )

        smb_query = search_wrapper.results.call_args_list[2].args[0]
        self.assertIn('"aerospace"', smb_query)
        self.assertIn('"small business" OR startup OR "under 500 employees"', smb_query)

    def test_fetch_jobs_smb_only_searches_direct_job_board_postings(self):
        lever_posting = {
            "title": "Lead Satellite Systems Engineer - Kepler",
            "link": "https://jobs.lever.co/kepler/223eeef2-f0e0-4656-a6d0-538074670f6d",
            "snippet": "Kepler Communications is hiring a Lead Satellite Systems Engineer.",
        }
        job_bank_posting = {
            "title": "Aerospace Engineer - Halifax - Job posting - Job Bank",
            "link": "https://www.jobbank.gc.ca/jobsearch/jobposting/48212273",
            "snippet": "Aerospace Engineer position in Canada.",
        }
        search_wrapper = Mock()
        search_wrapper.results.side_effect = [
            [],
            [],
            [],
            [lever_posting],
            [],
            [job_bank_posting],
        ]
        with patch(
            "langchain_community.utilities.DuckDuckGoSearchAPIWrapper",
            return_value=search_wrapper,
        ):
            jobs = agent.fetch_jobs(
                ["aerospace engineer"],
                "Canada",
                5,
                "aerospace",
                smb_only=True,
            )

        self.assertEqual(
            {job.url for job in jobs},
            {
                lever_posting["link"],
                job_bank_posting["link"],
            },
        )
        queries = [call.args[0] for call in search_wrapper.results.call_args_list]
        self.assertIn(
            "site:jobs.lever.co aerospace engineer Canada aerospace",
            queries,
        )
        self.assertIn(
            "site:jobbank.gc.ca/jobsearch/jobposting "
            "aerospace engineer Canada aerospace",
            queries,
        )

    def test_parser_exposes_smb_only_option(self):
        args = agent.build_parser().parse_args(
            ["--resume", "resume.pdf", "--industry", "aerospace", "--smb-focus"]
        )
        self.assertTrue(args.smb_only)

    def test_assess_posting_includes_target_industry(self):
        posting = agent.Posting("1", "Analyst", "Acme", "Remote", "Analyze data", "https://example.com")
        response = {
            "is_individual_posting": True,
            "industry_match": True,
            "location_match": "match",
            "posting_location": "Remote",
            "requirements": [
                {
                    "criterion": "Analysis",
                    "priority": "preferred",
                    "status": "met",
                    "evidence": "Analyzed data",
                }
            ],
            "summary": "Role aligns with experience",
        }
        with patch.object(
            agent, "_ollama_generate", return_value=json.dumps(response)
        ) as generate:
            agent.assess_posting(
                "resume", posting, "llama3.2:3b", "http://localhost:11434", "healthcare"
            )
        self.assertIn("The target industry is healthcare", generate.call_args.args[0])
        self.assertIn("do not treat search/category pages", generate.call_args.args[0])
        self.assertIn("https://example.com", generate.call_args.args[0])

    def test_assess_posting_prompt_compares_each_requested_location_level(self):
        posting = agent.Posting(
            "1",
            "Engineer",
            "Acme",
            "Toronto, Ontario, Canada",
            "Engineer in Ottawa, Ontario, Canada.",
            "https://example.com/job",
        )
        response = {
            "is_individual_posting": True,
            "industry_match": True,
            "location_match": "mismatch",
            "posting_location": "Ottawa, Ontario, Canada",
            "requirements": [
                {
                    "criterion": "Engineering",
                    "priority": "preferred",
                    "status": "met",
                    "evidence": "Engineering experience.",
                }
            ],
            "summary": "Ottawa is not the requested city, Toronto.",
        }
        with patch.object(
            agent, "_ollama_generate", return_value=json.dumps(response)
        ) as generate:
            assessment = agent.assess_posting(
                "Engineering experience.",
                posting,
                "llama3.2:3b",
                "http://localhost:11434",
            )

        prompt = generate.call_args.args[0]
        self.assertIn(
            "Every supplied component must match",
            prompt,
        )
        self.assertIn(
            "city in the same province is a mismatch",
            prompt,
        )
        self.assertFalse(assessment.qualified)
        self.assertEqual(assessment.posting.location, "Ottawa, Ontario, Canada")

    def test_assess_posting_identifies_and_rejects_explicit_location_mismatch(self):
        posting = agent.Posting(
            "1",
            "Research Engineer",
            "Waabi",
            "Canada",
            "Research Engineer. Boston, MA, USA. Experience with robotics.",
            "https://jobs.example/research-engineer",
        )
        response = {
            "is_individual_posting": True,
            "industry_match": True,
            "location_match": "mismatch",
            "posting_location": "Boston, MA, USA",
            "employer_name": "Waabi",
            "requirements": [
                {
                    "criterion": "Robotics experience",
                    "priority": "preferred",
                    "status": "met",
                    "evidence": "Worked with robotics.",
                }
            ],
            "summary": "The role is based in Boston, not Canada.",
        }
        with patch.object(
            agent, "_ollama_generate", return_value=json.dumps(response)
        ) as generate:
            assessment = agent.assess_posting(
                "Worked with robotics.",
                posting,
                "llama3.2:3b",
                "http://localhost:11434",
            )

        self.assertFalse(assessment.qualified)
        self.assertEqual(assessment.location_match, "mismatch")
        self.assertEqual(assessment.posting.location, "Boston, MA, USA")
        self.assertIn("Requested search location: Canada", generate.call_args.args[0])

    def test_assess_posting_marks_unknown_actual_location(self):
        posting = agent.Posting(
            "1", "Analyst", "Acme", "Canada", "Analyze data.", "https://example.com/job"
        )
        response = {
            "is_individual_posting": True,
            "industry_match": True,
            "location_match": "unclear",
            "posting_location": None,
            "requirements": [
                {
                    "criterion": "Data analysis",
                    "priority": "preferred",
                    "status": "met",
                    "evidence": "Analyzed datasets.",
                }
            ],
            "summary": "The location isn't specified in the source text.",
        }
        with patch.object(
            agent, "_ollama_generate", return_value=json.dumps(response)
        ):
            assessment = agent.assess_posting(
                "Analyzed datasets.",
                posting,
                "llama3.2:3b",
                "http://localhost:11434",
            )
        self.assertEqual(assessment.posting.location, "Not verified")
        self.assertEqual(assessment.location_match, "unclear")

    def test_assess_posting_calculates_weighted_requirements_score(self):
        posting = agent.Posting(
            "1", "Analyst", "Acme", "Remote", "SQL and reporting", "https://example.com/job"
        )
        response = {
            "is_individual_posting": True,
            "industry_match": True,
            "location_match": "match",
            "posting_location": "Remote",
            "requirements": [
                {
                    "criterion": "SQL",
                    "priority": "required",
                    "status": "met",
                    "evidence": "Used SQL daily",
                },
                {
                    "criterion": "Two years of analytics",
                    "priority": "required",
                    "status": "missing",
                    "evidence": "",
                },
                {
                    "criterion": "Dashboard experience",
                    "priority": "preferred",
                    "status": "met",
                    "evidence": "Built dashboards",
                },
            ],
            "summary": "SQL evidence is present, but the experience requirement is not.",
        }
        with patch.object(
            agent, "_ollama_generate", return_value=json.dumps(response)
        ) as generate:
            assessment = agent.assess_posting(
                "Used SQL daily and built dashboards",
                posting,
                "llama3.2:3b",
                "http://localhost:11434",
                "any",
            )

        self.assertEqual(assessment.required_score, 50)
        self.assertEqual(assessment.preferred_score, 100)
        self.assertEqual(assessment.score, 65)
        self.assertFalse(assessment.qualified)
        self.assertEqual(assessment.matching_skills, ["SQL", "Dashboard experience"])
        self.assertEqual(assessment.gaps, ["Two years of analytics"])
        self.assertIn("Python calculates it", generate.call_args.args[0])

    def test_assess_posting_has_no_score_without_assessable_requirements(self):
        posting = agent.Posting(
            "1", "Analyst", "Acme", "Remote", "", "https://example.com/job"
        )
        response = {
            "is_individual_posting": True,
            "industry_match": True,
            "location_match": "match",
            "posting_location": "Remote",
            "requirements": [],
            "summary": "The snippet is too short.",
        }
        with patch.object(
            agent, "_ollama_generate", return_value=json.dumps(response)
        ):
            assessment = agent.assess_posting(
                "resume", posting, "llama3.2:3b", "http://localhost:11434"
            )
        self.assertIsNone(assessment.score)
        self.assertFalse(assessment.qualified)
        self.assertIn("not contain enough assessable requirements", assessment.reason)

    def test_assess_posting_any_industry_does_not_penalize_sector(self):
        posting = agent.Posting(
            "1", "Analyst", "Acme", "Remote", "Analyze data", "https://example.com"
        )
        response = {
            "is_individual_posting": True,
            "industry_match": True,
            "location_match": "match",
            "posting_location": "Remote",
            "requirements": [
                {
                    "criterion": "Analysis",
                    "priority": "preferred",
                    "status": "met",
                    "evidence": "Analyzed data",
                }
            ],
            "summary": "Role aligns with experience",
        }
        with patch.object(
            agent, "_ollama_generate", return_value=json.dumps(response)
        ) as generate:
            agent.assess_posting(
                "resume", posting, "llama3.2:3b", "http://localhost:11434", "any"
            )
        prompt = generate.call_args.args[0]
        self.assertIn("candidate is open to any industry", prompt)
        self.assertIn("Do not penalize or reject a posting based on its industry", prompt)

    def test_generate_tailored_resumes_creates_docs_for_qualified_jobs_only(self):
        qualified = agent.Assessment(
            agent.Posting(
                "1",
                "Data Analyst",
                "Acme",
                "Remote",
                "Analyze datasets and communicate findings.",
                "https://example.com/jobs/analyst",
            ),
            85,
            True,
            "Strong match",
            ["SQL"],
            [],
        )
        below_threshold = agent.Assessment(
            agent.Posting(
                "2",
                "Engineer",
                "Other",
                "Remote",
                "Engineering role",
                "https://example.org/jobs/engineer",
            ),
            95,
            False,
            "Not qualified",
            [],
            [],
        )
        tailored_content = "# Summary\nExperienced analyst.\n\n# Skills\n- SQL"
        with tempfile.TemporaryDirectory() as directory:
            output_dir = Path(directory) / "tailored"
            with patch.object(
                agent,
                "_ollama_generate",
                return_value=json.dumps({"tailored_resume": tailored_content}),
            ) as generate:
                outputs = agent.generate_tailored_resumes(
                    "Source resume: analyst with SQL",
                    [qualified, below_threshold],
                    output_dir,
                    60,
                    "llama3.2:3b",
                    "http://localhost:11434",
                    "any",
                )

            self.assertEqual(len(outputs), 1)
            self.assertEqual(len(list(output_dir.glob("*.docx"))), 1)
            document = Document(outputs[0])
            content = "\n".join(paragraph.text for paragraph in document.paragraphs)
            self.assertIn("Tailored Resume Draft", content)
            self.assertIn("AI-assisted draft", content)
            self.assertIn("Target: Data Analyst | Acme", content)
            self.assertIn("Experienced analyst.", content)
            self.assertEqual(generate.call_count, 1)
            prompt = generate.call_args.args[0]
            self.assertIn("do not add or imply experience", prompt)
            self.assertIn("https://example.com/jobs/analyst", prompt)

    def test_generate_tailored_resumes_does_not_create_empty_output_dir(self):
        with tempfile.TemporaryDirectory() as directory:
            output_dir = Path(directory) / "tailored"
            self.assertEqual(
                agent.generate_tailored_resumes(
                    "resume",
                    [],
                    output_dir,
                    60,
                    "llama3.2:3b",
                    "http://localhost:11434",
                ),
                [],
            )
            self.assertFalse(output_dir.exists())

    def test_run_needs_no_external_search_api_key(self):
        with tempfile.TemporaryDirectory() as directory:
            resume = Path(directory) / "resume.txt"
            resume.write_text("Resume", encoding="utf-8")
            output_docx = Path(directory) / "jobs.docx"
            output_csv = Path(directory) / "jobs.csv"
            args = Namespace(
                resume=str(resume),
                max_results=15,
                min_score=60,
                model="llama3.2:3b",
                ollama_host="http://localhost:11434",
                location="Remote",
                output=str(output_docx),
                csv_output=str(output_csv),
            )
            with patch.dict("os.environ", {}, clear=True):
                with (
                    patch.object(agent, "ensure_ollama_running"),
                    patch.object(agent, "extract_search_queries", return_value=["analyst"]),
                    patch.object(agent, "fetch_jobs", return_value=[]),
                    patch.object(agent, "verify_posting_links", return_value=[]),
                ):
                    self.assertEqual(agent.run(args), 0)
            self.assertTrue(output_docx.is_file())
            self.assertTrue(output_csv.is_file())

    def test_run_excludes_explicitly_mismatched_locations_from_reports(self):
        with tempfile.TemporaryDirectory() as directory:
            resume = Path(directory) / "resume.txt"
            resume.write_text("Resume", encoding="utf-8")
            output_docx = Path(directory) / "jobs.docx"
            output_csv = Path(directory) / "jobs.csv"
            posting = agent.Posting(
                "1", "Research Engineer", "Waabi", "Canada", "Boston, MA, USA", "https://job.example"
            )
            mismatch = agent.Assessment(
                posting,
                100,
                False,
                "Posting location does not match Canada.",
                ["Research"],
                [],
                location_match="mismatch",
            )
            args = Namespace(
                resume=str(resume),
                max_results=5,
                min_score=60,
                model="llama3.2:3b",
                ollama_host="http://localhost:11434",
                location="Canada",
                industry=None,
                smb_only=False,
                output=str(output_docx),
                csv_output=str(output_csv),
            )
            with (
                patch.object(agent, "ensure_ollama_running"),
                patch.object(agent, "extract_search_queries", return_value=["engineer"]),
                patch.object(agent, "fetch_jobs", return_value=[posting]),
                patch.object(agent, "verify_posting_links", return_value=[posting]),
                patch.object(agent, "assess_posting", return_value=mismatch),
            ):
                self.assertEqual(agent.run(args), 0)

            with output_csv.open(encoding="utf-8-sig", newline="") as file:
                rows = list(csv.DictReader(file))
            self.assertEqual(rows, [])

    def test_run_smb_focus_retains_unverified_size_candidates(self):
        with tempfile.TemporaryDirectory() as directory:
            resume = Path(directory) / "resume.txt"
            resume.write_text("Resume", encoding="utf-8")
            output_docx = Path(directory) / "jobs.docx"
            output_csv = Path(directory) / "jobs.csv"
            postings = [
                agent.Posting("1", "Engineer", "SmallCo", "Canada", "", "https://small.example/job"),
                agent.Posting("2", "Engineer", "BigCo", "Canada", "", "https://big.example/job"),
                agent.Posting("3", "Engineer", "UnknownCo", "Canada", "", "https://unknown.example/job"),
            ]
            assessments = [
                agent.Assessment(posting, 80, True, "Match", [], [])
                for posting in postings
            ]
            verified_assessments = [
                agent.dataclass_replace(
                    assessments[0],
                    posting=agent.dataclass_replace(
                        assessments[0].posting,
                        company_size="51-200 employees",
                        company_size_evidence="SmallCo has 51-200 employees.",
                    ),
                ),
                agent.dataclass_replace(
                    assessments[1],
                    posting=agent.dataclass_replace(
                        assessments[1].posting,
                        company_size="1,001-5,000 employees",
                        company_size_evidence="BigCo has 1,001-5,000 employees.",
                    ),
                ),
                assessments[2],
            ]
            args = Namespace(
                resume=str(resume),
                max_results=5,
                min_score=60,
                model="llama3.2:3b",
                ollama_host="http://localhost:11434",
                location="Canada",
                industry="aerospace",
                smb_only=True,
                output=str(output_docx),
                csv_output=str(output_csv),
            )
            with (
                patch.object(agent, "ensure_ollama_running"),
                patch.object(agent, "extract_search_queries", return_value=["engineer"]),
                patch.object(agent, "fetch_jobs", return_value=postings) as fetch_jobs,
                patch.object(agent, "verify_posting_links", return_value=postings),
                patch.object(agent, "assess_posting", side_effect=assessments),
                patch.object(
                    agent,
                    "verify_employer_facts",
                    return_value=verified_assessments,
                ),
            ):
                self.assertEqual(agent.run(args), 0)

            fetch_jobs.assert_called_once_with(
                ["engineer"], "Canada", 5, "aerospace", True
            )
            with output_csv.open(encoding="utf-8", newline="") as file:
                rows = list(csv.DictReader(file))
            self.assertEqual(
                [row["company"] for row in rows], ["SmallCo", "UnknownCo"]
            )
            self.assertEqual(rows[1]["company_size"], "Not verified")

    def test_ensure_ollama_running_does_not_start_duplicate_server(self):
        with (
            patch.object(agent, "_ollama_api_available", return_value=True),
            patch.object(agent.subprocess, "Popen") as popen,
        ):
            agent.ensure_ollama_running("http://localhost:11434")
        popen.assert_not_called()

    def test_ensure_ollama_running_starts_local_server_and_waits_for_api(self):
        process = Mock()
        process.poll.return_value = None
        with (
            patch.object(
                agent, "_ollama_api_available", side_effect=[False, False, True]
            ),
            patch.object(agent.shutil, "which", return_value="/usr/local/bin/ollama"),
            patch.object(agent.subprocess, "Popen", return_value=process) as popen,
            patch.object(agent.time, "sleep"),
        ):
            agent.ensure_ollama_running("http://localhost:11434", timeout=5)

        self.assertEqual(
            popen.call_args.args[0], ["/usr/local/bin/ollama", "serve"]
        )
        self.assertEqual(popen.call_args.kwargs["stdin"], agent.subprocess.DEVNULL)
        self.assertTrue(popen.call_args.kwargs["start_new_session"])

    def test_ensure_ollama_running_reports_missing_executable(self):
        with (
            patch.object(agent, "_ollama_api_available", return_value=False),
            patch.object(agent.shutil, "which", return_value=None),
        ):
            with self.assertRaisesRegex(RuntimeError, "command was not found"):
                agent.ensure_ollama_running("http://localhost:11434")

    def test_ensure_ollama_running_does_not_start_remote_server(self):
        with (
            patch.object(agent, "_ollama_api_available", return_value=False),
            patch.object(agent.shutil, "which") as which,
        ):
            with self.assertRaisesRegex(RuntimeError, "only auto-starts Ollama"):
                agent.ensure_ollama_running("http://ollama.example:11434")
        which.assert_not_called()

    def test_parser_exposes_opt_in_resume_tailoring(self):
        args = agent.build_parser().parse_args(
            ["--resume", "resume.txt", "--tailor-resumes"]
        )
        self.assertTrue(args.tailor_resumes)
        self.assertEqual(args.tailored_output_dir, "tailored_resumes")

    def test_ollama_request_uses_local_generate_endpoint_and_json_mode(self):
        payload = {"response": '{"ok": true}'}

        class FakeResponse:
            def __enter__(self):
                return self

            def __exit__(self, *_args):
                return None

            def read(self):
                return json.dumps(payload).encode()

        with patch.object(agent, "urlopen", return_value=FakeResponse()) as urlopen:
            result = agent._ollama_generate(
                "return JSON", "qwen2.5:3b", "http://localhost:11434"
            )
        request = urlopen.call_args.args[0]
        self.assertEqual(request.full_url, "http://localhost:11434/api/generate")
        request_data = json.loads(request.data)
        self.assertEqual(request_data["model"], "qwen2.5:3b")
        self.assertEqual(request_data["prompt"], "return JSON")
        self.assertEqual(request_data["format"], "json")
        self.assertFalse(request_data["stream"])
        self.assertEqual(result, '{"ok": true}')
        self.assertEqual(
            urlopen.call_args.kwargs["timeout"],
            agent.DEFAULT_OLLAMA_GENERATE_TIMEOUT,
        )

    def test_ollama_generate_uses_configured_timeout(self):
        payload = {"response": '{"ok": true}'}

        class FakeResponse:
            def __enter__(self):
                return self

            def __exit__(self, *_args):
                return None

            def read(self):
                return json.dumps(payload).encode()

        with (
            patch.dict("os.environ", {"OLLAMA_GENERATE_TIMEOUT": "900"}),
            patch.object(agent, "urlopen", return_value=FakeResponse()) as urlopen,
        ):
            agent._ollama_generate(
                "return JSON", "llama3.2:3b", "http://localhost:11434"
            )

        self.assertEqual(urlopen.call_args.kwargs["timeout"], 900)

    def test_ollama_generate_reports_generation_timeout_and_recovery(self):
        with patch.object(agent, "urlopen", side_effect=TimeoutError("timed out")):
            with self.assertRaisesRegex(
                RuntimeError,
                "Ollama generation timed out after 600 seconds.*Increase "
                "OLLAMA_GENERATE_TIMEOUT",
            ):
                agent._ollama_generate(
                    "long prompt", "llama3.2:3b", "http://localhost:11434"
                )

    def test_ollama_generation_timeout_must_be_positive(self):
        with patch.dict("os.environ", {"OLLAMA_GENERATE_TIMEOUT": "0"}):
            with self.assertRaisesRegex(
                RuntimeError, "OLLAMA_GENERATE_TIMEOUT must be a positive"
            ):
                agent._ollama_generate_timeout()


if __name__ == "__main__":
    unittest.main()
